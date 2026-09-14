"""PDF / DOCX / plain-text extraction."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from app.services.documents import (
    DocumentError,
    ParsedDocument,
    normalise_text,
    parse_document,
    parse_pasted_text,
)

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"
CLAUSE = (
    "The agreement renews automatically on 15 November 2026 for a further period of "
    "twelve (12) months unless cancelled at least 30 days before renewal."
)


def _make_pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    y = 800
    for line in text.split("\n"):
        pdf.drawString(50, y, line[:95])
        y -= 14
    pdf.save()
    return buffer.getvalue()


def _make_docx(text: str) -> bytes:
    import docx

    document = docx.Document()
    for line in text.split("\n"):
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_plain_text_upload():
    parsed = parse_document(CLAUSE.encode(), "clause.txt", "text/plain")
    assert isinstance(parsed, ParsedDocument)
    assert parsed.kind == "text"
    assert "30 days before renewal" in parsed.text


def test_pdf_upload_roundtrip():
    parsed = parse_document(_make_pdf(CLAUSE), "contract.pdf", "application/pdf")
    assert parsed.kind == "pdf"
    assert "renews automatically" in parsed.text


def test_docx_upload_roundtrip():
    data = _make_docx(CLAUSE)
    parsed = parse_document(
        data,
        "contract.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert parsed.kind == "docx"
    assert "30 days before renewal" in parsed.text


def test_docx_tables_are_included():
    import docx

    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Notice period"
    table.rows[0].cells[1].text = "60 days"
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = parse_document(buffer.getvalue(), "terms.docx")
    assert "Notice period | 60 days" in parsed.text


def test_pdf_detected_by_magic_bytes_without_extension():
    parsed = parse_document(_make_pdf(CLAUSE), "no-extension", "application/octet-stream")
    assert parsed.kind == "pdf"


def test_corrupt_pdf_raises_document_error():
    with pytest.raises(DocumentError):
        parse_document(b"%PDF-1.4 this is not really a pdf", "broken.pdf", "application/pdf")


def test_corrupt_docx_raises_document_error():
    with pytest.raises(DocumentError):
        parse_document(b"PK\x03\x04 not a docx", "broken.docx")


def test_empty_upload_raises():
    with pytest.raises(DocumentError):
        parse_document(b"", "empty.txt")


def test_oversized_upload_raises():
    with pytest.raises(DocumentError):
        parse_document(b"x" * (10 * 1024 * 1024 + 1), "huge.txt", "text/plain")


def test_unsupported_type_raises():
    with pytest.raises(DocumentError):
        parse_document(b"\x89PNG\r\n", "scan.png", "image/png")


def test_blank_paste_raises():
    with pytest.raises(DocumentError):
        parse_pasted_text("   \n\n  ")


def test_paste_is_normalised():
    parsed = parse_pasted_text("  Renews   on\t15 November 2026 \n\n\n\n Notice: 30 days ")
    assert parsed.text == "Renews on 15 November 2026\n\nNotice: 30 days"


def test_normalise_handles_windows_line_endings():
    assert normalise_text("a\r\nb\r\n") == "a\nb"


@pytest.mark.parametrize(
    "name", ["saas-contract.txt", "insurance-renewal.txt", "service-agreement.txt"]
)
def test_bundled_examples_parse(name):
    parsed = parse_document((EXAMPLES / name).read_bytes(), name, "text/plain")
    assert len(parsed.text) > 200
